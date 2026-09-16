import io

from docx import Document as DocxDocument
from openpyxl import load_workbook
import pdfplumber

CHUNK_SIZE = 1000
CHUNK_OVERLAP = 100


def extract_text_and_chunk(file_bytes: bytes, file_type: str) -> list[tuple[str, int]]:
    if file_type == "PDF":
        pages = _extract_pdf(file_bytes)
    elif file_type == "DOCX":
        pages = _extract_docx(file_bytes)
    elif file_type == "XLSX":
        pages = _extract_xlsx(file_bytes)
    else:
        raise ValueError(f"Unsupported file type: {file_type}")

    chunks: list[tuple[str, int]] = []
    for page_number, page_text in pages:
        for chunk_text in _split_into_chunks(page_text):
            chunks.append((chunk_text, page_number))

    return chunks


def _extract_pdf(file_bytes: bytes) -> list[tuple[int, str]]:
    pages: list[tuple[int, str]] = []

    with pdfplumber.open(io.BytesIO(file_bytes)) as pdf:
        for page_number, page in enumerate(pdf.pages, start=1):
            parts = []

            text = page.extract_text() or ""
            if text.strip():
                parts.append(text)

            tables = page.extract_tables()
            for table in tables:
                for row in table:
                    row_text = " | ".join(str(cell).strip() if cell else "" for cell in row)
                    if row_text.strip(" |"):
                        parts.append(row_text)

            pages.append((page_number, "\n".join(parts)))

    return pages


def _extract_docx(file_bytes: bytes) -> list[tuple[int, str]]:
    doc = DocxDocument(io.BytesIO(file_bytes))

    parts = [paragraph.text for paragraph in doc.paragraphs if paragraph.text.strip()]

    for table in doc.tables:
        for row in table.rows:
            row_text = " | ".join(cell.text.strip() for cell in row.cells)
            if row_text.strip(" |"):
                parts.append(row_text)

    full_text = "\n".join(parts)
    return [(1, full_text)]


def _extract_xlsx(file_bytes: bytes) -> list[tuple[int, str]]:
    workbook = load_workbook(io.BytesIO(file_bytes), data_only=True)
    pages: list[tuple[int, str]] = []

    for sheet_index, sheet in enumerate(workbook.worksheets, start=1):
        rows_text = []
        for row in sheet.iter_rows(values_only=True):
            row_text = " | ".join(str(cell) for cell in row if cell is not None)
            if row_text:
                rows_text.append(row_text)
        pages.append((sheet_index, "\n".join(rows_text)))

    return pages


def _split_into_chunks(text: str) -> list[str]:
    text = text.strip()
    if not text:
        return []

    chunks = []
    start = 0
    while start < len(text):
        end = start + CHUNK_SIZE
        chunks.append(text[start:end])
        start = end - CHUNK_OVERLAP

    return chunks